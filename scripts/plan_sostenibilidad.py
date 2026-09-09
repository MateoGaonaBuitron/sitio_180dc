"""Adaptador del plan semanal de agosto 2026-II, conservando el Excel original.

La deuda heredada es un pasivo pendiente, no un pago de agosto. El saldo del
plan se concilia con disponible menos deuda. No se importan meses futuros.
"""
from decimal import Decimal

import openpyxl
import pandas as pd

from helpers import InformeMes


def cargar_agosto_plan(path, narrativa):
    with open(path, "rb") as source:
        book = openpyxl.load_workbook(source, data_only=True)
    try:
        sheet = book["Flujo de Caja"]

        def dinero(row, col, *, required=False):
            value = sheet.cell(row, col).value
            if value is None and not required:
                return Decimal("0")
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(f"Monto ausente o inválido en {sheet.cell(row, col).coordinate}")
            result = Decimal(str(value))
            if not result.is_finite():
                raise ValueError("El plan contiene un monto no finito")
            return result.quantize(Decimal("0.01"))

        # Validar la estructura conocida para evitar leer otras filas por error.
        for coordinate, expected in {
            "C5": "S1", "D5": "S2", "E5": "S3",
            "C6": "17/08", "D6": "24/08", "E6": "31/08",
            "A7": "SALDO INICIAL", "A24": "Total Ingresos",
            "A27": "DEUDA", "A42": "Total Egresos", "A45": "SALDO FINAL",
        }.items():
            if sheet[coordinate].value != expected:
                raise ValueError(f"Cambió el formato del plan en {coordinate}; revisar el adaptador")

        inicial = dinero(7, 3, required=True)
        disponible, deuda = inicial, Decimal("0")
        rows, semanas, fuentes, egresos = [], [], [], []
        for col in (3, 4, 5):
            semana = sheet.cell(5, col).value
            fecha = f"Semana del {sheet.cell(6, col).value}/2026"
            ingresos, gastos = Decimal("0"), Decimal("0")
            deuda_semana = dinero(27, col)
            if deuda_semana < 0:
                raise ValueError("Revisar el tratamiento de la deuda: se detectó una reducción")
            for tipo, rango in (("Ingreso", range(9, 24)), ("Egreso", range(27, 42))):
                for row in rango:
                    if row == 27:
                        continue
                    monto = dinero(row, col)
                    if not monto:
                        continue
                    categoria = sheet.cell(row, 1).value
                    concepto = sheet.cell(row, 2).value
                    if monto < 0 or not categoria or not concepto:
                        raise ValueError(f"Revisar concepto o monto de la fila {row}")
                    detalle = f"{concepto}. Registro semanal {semana}; el Excel no especifica el día del movimiento."
                    rows.append({"Tipo": tipo, "Monto": float(monto), "Categoria": categoria,
                                 "Descripcion": concepto, "Fecha": fecha, "EsInversion": False})
                    if tipo == "Ingreso":
                        ingresos += monto
                        fuentes.append({"nombre": categoria, "monto": float(monto),
                                        "fecha": fecha, "detalle": detalle, "categoria_ui": categoria})
                    else:
                        gastos += monto
                        egresos.append({"titulo": concepto, "monto": float(monto), "fecha": fecha,
                                        "para_que": detalle, "categoria_ui": categoria, "es_inversion": False})
            esperado_inicial = disponible - deuda
            deuda += deuda_semana
            disponible += ingresos - gastos
            neto = disponible - deuda
            for actual, expected, label in (
                (dinero(7, col, required=True), esperado_inicial, "saldo inicial"),
                (dinero(24, col, required=True), ingresos, "ingresos"),
                (dinero(42, col, required=True), gastos + deuda_semana, "egresos del plan"),
                (dinero(45, col, required=True), neto, "saldo final del plan"),
            ):
                if actual != expected:
                    raise ValueError(f"No cuadra {label} en {semana}: Excel {actual}, calculado {expected}. Recalcula y guarda el Excel.")
            semanas.append({"semana": semana, "fecha": fecha, "ingresos": float(ingresos),
                            "gastos": float(gastos), "disponible": float(disponible),
                            "deuda": float(deuda), "neto": float(neto)})

        frame = pd.DataFrame(rows, columns=["Tipo", "Monto", "Categoria", "Descripcion", "Fecha", "EsInversion"])
        entradas = frame[frame["Tipo"] == "Ingreso"].copy()
        salidas = frame[frame["Tipo"] == "Egreso"].copy()
        info = InformeMes(
            mes="Agosto", ciclo="2026-2", responsable="Área financiera",
            actualizacion=narrativa["actualizacion"], saldo_inicial=float(inicial),
            ingresos=float(entradas["Monto"].sum()), egresos_op=float(salidas["Monto"].sum()),
            inversion=0.0, saldo_final=float(disponible), df=frame,
            df_ingresos=entradas, df_egresos=salidas, df_inversion=frame.iloc[:0].copy(),
        )
        contenido = {**narrativa, "fuentes_ingresos": fuentes, "egresos": egresos,
                     "semanas": semanas, "deuda": {**narrativa["deuda"], "monto": float(deuda)}}
        return info, contenido
    finally:
        book.close()
